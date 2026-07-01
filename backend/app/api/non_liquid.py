import io

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user
from app.core.database import get_db
from app.models.non_liquid import NonLiquid
from app.models.product import Product
from app.models.user import User
from app.schemas.non_liquid import NonLiquidOut, NonLiquidUploadResult

router = APIRouter(prefix="/non-liquid", tags=["non_liquid"])


@router.get("", response_model=list[NonLiquidOut])
async def list_non_liquid(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[NonLiquid]:
    stmt = (
        select(NonLiquid)
        .where(NonLiquid.user_id == current_user.id)
        .order_by(NonLiquid.id.desc())
    )
    return list(await db.scalars(stmt))


@router.post("/upload", response_model=NonLiquidUploadResult)
async def upload_excel(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> NonLiquidUploadResult:
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only .xlsx/.xlsm files are supported",
        )
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read Excel file: {exc}",
        ) from exc

    ws = wb.active
    existing_oems = set(
        await db.scalars(
            select(Product.oem_number).where(Product.user_id == current_user.id)
        )
    )

    total = created = duplicates = 0
    header_skipped = False
    for row in ws.iter_rows(values_only=True):
        if not header_skipped:
            header_skipped = True
            continue
        if not row or all(c is None for c in row):
            continue
        total += 1
        oem = str(row[0]).strip() if row[0] is not None else None
        article = str(row[1]).strip() if len(row) > 1 and row[1] else None
        name = str(row[2]).strip() if len(row) > 2 and row[2] else None
        brand = str(row[3]).strip() if len(row) > 3 and row[3] else None
        if oem and oem in existing_oems:
            duplicates += 1
            continue
        db.add(
            NonLiquid(
                user_id=current_user.id,
                oem_number=oem,
                article=article,
                name=name,
                brand=brand,
                is_checked=bool(oem and oem not in existing_oems),
            )
        )
        created += 1

    await db.commit()
    return NonLiquidUploadResult(
        total_rows=total, created=created, duplicates=duplicates
    )
